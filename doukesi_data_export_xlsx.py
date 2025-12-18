import pandas as pd
from dousike_data_class import G_stats, G_orders, G_redeems
from dousike_data_meituan_class import G_meituan, MeituanOrder

import re

def normalize_name(name: str) -> str:
    if not name:
        return ""
    name = name.lower()
    name = re.sub(r"[（）()·•\s\-]", "", name)
    name = re.sub(r"(店|门店|分店)$", "", name)
    return name


df_mt = pd.DataFrame([
    {
        "date": o.pay_time.date(),
        "store_id": str(o.poi_id),
        "store_name": o.poi_name,
        "meituan_order_cnt": 1
    }
    for o in G_meituan
])

df_mt["date"] = pd.to_datetime(df_mt["date"])

df_mt_daily = (
    df_mt
    .groupby(["store_id", "store_name", "date"], as_index=False)
    .agg({"meituan_order_cnt": "sum"})
)


def stats_to_dataframe(stats_dict):
    """
    stats_dict: dict[(store_id, stat_date), StoreDailyStat]
    """
    data = []
    for stat in stats_dict.values():
        data.append({
            "store_id": stat.store_id,
            "store_name": stat.store_name,
            "date": stat.stat_date,
            "coupon_order_cnt": stat.order_cnt,
            "order_with_lday_cnt": max(0, stat.order_cnt - stat.redeem_from_prev_days),
            "redeem_cnt": stat.redeem_cnt,
            "redeem_from_prev_days": stat.redeem_from_prev_days
        })
    df = pd.DataFrame(data)
    df["date"] = pd.to_datetime(df["date"])
    return df

# 假设 stats_dict 已经生成
# stats_dict = build_store_daily_stats(orders, redeems)
df_coupon = stats_to_dataframe(G_stats)


df_coupon["date"] = pd.to_datetime(df_coupon["date"])
df_coupon["store_id"] = df_coupon["store_id"].astype(str)





df_mt_daily["norm_name"] = df_mt_daily["store_name"].apply(normalize_name)
df_coupon["norm_name"] = df_coupon["store_name"].apply(normalize_name)



from difflib import SequenceMatcher
import pandas as pd

def similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()

def build_store_mapping(
    df_mt_daily: pd.DataFrame,
    df_coupon: pd.DataFrame,
    threshold: float = 0.7
) -> pd.DataFrame:

    mt_stores = (
        df_mt_daily[["store_id", "store_name", "norm_name"]]
        .drop_duplicates()
        .reset_index(drop=True)
    )

    coupon_stores = (
        df_coupon[["store_id", "store_name", "norm_name"]]
        .drop_duplicates()
        .reset_index(drop=True)
    )

    mappings = []

    for _, mt_row in mt_stores.iterrows():
        best_score = 0
        best_row = None

        for _, cp_row in coupon_stores.iterrows():
            score = similarity(mt_row["norm_name"], cp_row["norm_name"])

            if score > best_score:
                best_score = score
                best_row = cp_row

        if best_row is None or best_score < threshold:
            continue

        mappings.append({
            "mt_store_id": mt_row["store_id"],
            "mt_store_name": mt_row["store_name"],
            "coupon_store_id": best_row["store_id"],
            "coupon_store_name": best_row["store_name"],
            "score": round(best_score, 3)
        })

    return pd.DataFrame(mappings)


df_store_mapping = build_store_mapping(df_mt_daily, df_coupon)

df_coupon_mapped = (
    df_coupon
    .merge(
        df_store_mapping[[
            "coupon_store_id",
            "mt_store_id",
            "mt_store_name"
        ]],
        left_on="store_id",
        right_on="coupon_store_id",
        how="left"
    )
)
df_coupon_mapped = (
    df_coupon_mapped
    .rename(columns={
        "mt_store_id": "store_id",
        "mt_store_name": "store_name"
    })
    .drop(columns=["coupon_store_id"])
)

print(df_coupon_mapped.columns.tolist())

cols = df_coupon_mapped.columns.tolist()
cols[0] = "mt_store_id"
cols[1] = "mt_store_name"
df_coupon_mapped.columns = cols
df_coupon_mapped[df_coupon_mapped["store_id"].isna()][
    ["store_name"]
].drop_duplicates()
df_store_mapping["score"].describe()
df_store_mapping.sort_values("score").head(10)


df_store_daily = (
    df_mt_daily
    .merge(
        df_coupon_mapped[
            [
                "store_id",
                "store_name",
                "date",
                "coupon_order_cnt",
                "redeem_cnt",
                "redeem_from_prev_days"
            ]
        ],
        on=["store_id", "store_name", "date"],
        how="left"
    )
    .fillna(0)
)



# df_store_daily = (
#     df_mt_daily
#     .merge(
#         df_coupon,
#         on=["store_id", "store_name", "date"],
#         how="left"
#     )
#     .fillna(0)
# )

df_store_daily["coupon_redeem_cnt"] = (
    df_store_daily["redeem_cnt"]
    - df_store_daily["redeem_from_prev_days"]
).clip(lower=0)

df_store_daily["inflow_rate"] = (
    df_store_daily["coupon_order_cnt"]
    / (df_store_daily["meituan_order_cnt"] + (df_store_daily["coupon_order_cnt"] - df_store_daily["coupon_redeem_cnt"]))
)

df_store_daily["inflow_rate"] = (
    df_store_daily["inflow_rate"]
    .replace([float("inf")], 0)
    .fillna(0)
)



# 假设 df_store_daily 已经准备好，并包含以下列：
# 'store_name', 'date', 'meituan_order_cnt', 'coupon_order_cnt', 'coupon_redeem_cnt', 'inflow_rate'

# ---- 1️⃣ 固定子行顺序 ----
metric_order = ["meituan_order_cnt", "coupon_order_cnt", "coupon_redeem_cnt", "inflow_rate"]
metric_map = {
    "meituan_order_cnt": "美团下单数",
    "coupon_order_cnt": "优惠券下单数",
    "coupon_redeem_cnt": "优惠券核销数",
    "inflow_rate": "引流率"
}

df_long = (
    df_store_daily[["store_name", "date"] + metric_order]
    .melt(
        id_vars=["store_name", "date"],
        value_vars=metric_order,
        var_name="metric",
        value_name="value"
    )
)

# 使用 category 设置顺序
df_long["metric"] = pd.Categorical(df_long["metric"], categories=metric_order, ordered=True)
df_long["metric_name"] = df_long["metric"].map(metric_map)

# ---- 2️⃣ Pivot 成表格 ----
df_export = (
    df_long.pivot_table(
        index=["store_name", "metric_name"],
        columns="date",
        values="value",
        aggfunc="sum"
    )
    .sort_index(level=0)
)


from openpyxl import load_workbook
from openpyxl.styles import PatternFill, numbers
from openpyxl.utils import get_column_letter
# ---- 3️⃣ 导出到 Excel ----
excel_path = "店铺每日指标.xlsx"
df_export.to_excel(excel_path, sheet_name="Sheet1")

# ---- 4️⃣ 打开 Excel 做格式化 ----
wb = load_workbook(excel_path)
ws = wb.active

# 日期列只显示 yyyy-mm-dd
for col in range(2, ws.max_column + 1):
    ws.cell(row=1, column=col).number_format = 'yyyy-mm-dd'

# 店铺大行交替颜色 & 冻结首列/首行
fill_colors = ["FFFFFF", "F2F2F2"]
store_names = df_export.index.get_level_values(0).unique()
row_cursor = 2  # 数据从第2行开始
for i, store in enumerate(store_names):
    fill = PatternFill(start_color=fill_colors[i % 2], end_color=fill_colors[i % 2], fill_type="solid")
    for r in range(4):  # 每个店铺4行子行
        for c in range(1, ws.max_column + 1):  # 包含首列 store_name
            cell = ws.cell(row=row_cursor + r, column=c)
            cell.fill = fill
            # 百分比列格式化
            if df_export.index.get_level_values(1)[row_cursor - 2 + r] == "引流率":
                cell.number_format = '0.00%'
            # 其他数字列千分位
            elif c > 1:
                cell.number_format = '#,##0'
    row_cursor += 4

# 冻结首列和首行
ws.freeze_panes = 'B2'

wb.save(excel_path)
print(f"导出完成: {excel_path}")

# # ---- 3️⃣ Streamlit 展示 ----
# # 将行按照店铺分组，然后每个店铺大行设置不同背景色
# def style_store_rows(row):
#     # 每个 store_name 都是一组，给背景色交替
#     store_names = df_export.index.get_level_values(0).unique()
#     color_map = {store: ("#f2f2f2" if i % 2 == 0 else "white") for i, store in enumerate(store_names)}
#     return ["background-color: {}".format(color_map[row.name[0]])]*len(row)
#
# st.dataframe(df_export.style.apply(style_store_rows, axis=1))