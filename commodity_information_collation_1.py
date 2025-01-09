import sys, os
import time
from dataclasses import dataclass, field
import json
import openpyxl
from thefuzz import fuzz

from address_recognition import address_rec

@dataclass(repr=False)
class commodity_info:
    id: str = field(default=None)
    count: int = field(default=1)
    total_count: int = field(default=0)
    type: int = field(default=0)
    name: str = field(default='')
    name_another: str = field(default='')
    price: float = field(default=None)
    total_price: float = field(default=None)

    geishitang: float = field(default=None)
    geiyuangong: float = field(default=None)

    def __post_init__(self):
        self.total_price = self.price * self.count

    def take(self):
        self.total_count = self.total_count - self.count

    def addcount(self, _count):
        self.total_count = self.total_count + 1 * _count

def softID(_it, _itl, _score):
    res_score = {}
    for it in _itl:
        ratio_it = fuzz.partial_ratio(it, _it)
        res_score[it] = ratio_it

    max_key_value = max(res_score.items(), key=lambda item: item[1])
    # print(sorted_dict)
    return max_key_value[0], max_key_value[1]

def loaddata(_sheet, type, id_idx=None):
    backdata = None
    global commodity_data
    for i in range(_sheet.max_row + 1):
        if i < 2:
            continue
        data_id = _sheet.cell(row=i, column=id_idx if id_idx else 1).value
        if data_id is None:
            continue
        ttype = type
        data_type = 1 if ttype else 0
        data_name = _sheet.cell(row=i, column=5).value
        data_price = _sheet.cell(row=i, column=11).value
        # data_tcount = sheet.cell(row=i, column=9).value
        data_tcount = 0

        st = _sheet.cell(row=i, column=11).value
        yg = _sheet.cell(row=i, column=9).value

        if data_name is None:
            continue
        data_count = 1
        # if '金龙鱼' in data_name or '维达' in data_name:
        #     data_count = 2
        if '金龙鱼' in data_name:
            data_count = 1
            data_price = int(data_price / 2)
            st = int(st / 2)
            yg = int(yg / 2)

        t_commodity = commodity_info(id=data_id, count=data_count, total_count=data_tcount, type=data_type, name=data_name, price=data_price, geishitang=st, geiyuangong=yg)
        commodity_data[data_name + str(data_type)] = t_commodity

commodity_data={}

def foo(commodity_data_path, src_data_path, resultname):


    workbook = openpyxl.load_workbook(commodity_data_path, data_only=True)
    # sheet = workbook['仓发1']
    # loaddata(sheet, 0)
    # sheet = workbook['代发1']
    # loaddata(sheet, 1)

    sheet = workbook['Sheet1']
    loaddata(sheet, 0, 3)
    sheet = workbook['Sheet1']
    loaddata(sheet, 1, 4)


    workbook = openpyxl.load_workbook(src_data_path, data_only=True)
    sheet = workbook['Sheet1']

    wb = openpyxl.Workbook()
    # 选择默认的工作表
    # wsheet = wb.active
    ws1 = wb.create_sheet('代发')
    ws2 = wb.create_sheet('直发')
    ws3 = wb.create_sheet('other')

    # 给工作表重命名
    # wsheet.title = 'data'

    # 最大行 sheet.max_row # 最大列 sheet.max_column
    backdata = None

    totalAcountprice = {}
    # titleNameList = ['订单号', '姓名', '电话', '商家SKU编码', '单价', '数量', '应付', '实付', '收货人', '收货手机号', '省', '市', '区', '地址', '买家留言'] # back
    titleNameList = ['商品编号', '商品价格', '商品数量', '买家应付货款', '收货人名称', '联系电话', '联系手机', '省', '市', '区', '街道', '地址', '发票抬头', '买家留言', '收货人应付邮费'] # 0311
    for i in range(sheet.max_row + 1):
        if i < 2:
            if i == 1:
                ws1.append(titleNameList)
                ws2.append(titleNameList)
                # ws3.append(['商品编号', '商品名称', '销售总数', '下单价', '给食堂', '给员工', '总金额'])
                ws3.append(['商品编号', '商品名称', '销售总数', '下单价', '下单价总金额', '员工价', '员工价总金额', '食堂价', '食堂价总金额'])
            continue
        data_id = sheet.cell(row=i, column=1).value
        data_name = sheet.cell(row=i, column=7).value
        data_ipone = sheet.cell(row=i, column=8).value
        data_d1 = sheet.cell(row=i, column=9).value
        data_d2 = sheet.cell(row=i, column=10).value
        data_d3 = sheet.cell(row=i, column=11).value

        addr_start_idx = 16
        while sheet.cell(row=i, column=addr_start_idx).value is not None:
            if '跳过' in sheet.cell(row=i, column=addr_start_idx).value:
                addr_start_idx = addr_start_idx + 4
            else:
                break
        data_sname = sheet.cell(row=i, column=addr_start_idx-3).value
        data_sipone = sheet.cell(row=i, column=addr_start_idx-2).value
        if data_sname is None:
            data_sname = data_name
        if data_sipone is None:
            data_sipone = data_ipone
        # data_saddr = '-'.join([sheet.cell(row=i, column=8).value, sheet.cell(row=i, column=9).value, sheet.cell(row=i, column=10).value])
        data_saddr = sheet.cell(row=i, column=addr_start_idx-1).value
        data_saddr_all = sheet.cell(row=i, column=addr_start_idx).value
        if data_saddr_all is None:
            data_saddr_all = '浙江省嘉兴市南湖区中环南路1385号中国人民银行嘉兴支行'
            data_saddr = '浙江-嘉兴市-南湖区'
        data_total_amount = sheet.cell(row=i, column=12).value
        if data_id is None:
            continue

        data_sku_z_list = []
        data_count_z_list = []
        data_price_z_list = []
        data_sku_c_list = []
        data_count_c_list = []
        data_price_c_list = []
        for tdata in [data_d1, data_d2, data_d3]:
            for it_di in tdata.split('+'):
                if it_di == '(空)':
                    continue
                # tcount = int(it_di.split('〖')[-1][0])
                tcount = 1
                resId, ratio = softID(it_di.split('〖')[0], commodity_data.keys(), 97)
                if ratio >= 90:
                    it = resId
                    print(' {} : {} -- {}'.format(ratio, it, it_di))
                    if commodity_data[it].type == 1:
                        data_sku_z_list.append(commodity_data[it].id)
                        data_count_z_list.append(tcount)
                        data_price_z_list.append(commodity_data[it].price)
                    else:
                        data_sku_c_list.append(commodity_data[it].id)
                        data_count_c_list.append(tcount)
                        data_price_c_list.append(commodity_data[it].price)
                    # commodity_data[it].take()
                    commodity_data[it].addcount(tcount)
                else:
                    print('!!!!!!!!!!!!!!!!!! : {}', it_di)
                    if '金龙鱼' in it_di:
                        tkey1 = '金龙鱼乳玉皇妃稻香贡米5kg0'
                        tkey2 = '金龙鱼葵花籽油5L0'
                        if commodity_data[tkey1].type == 1:
                            data_sku_z_list.append(commodity_data[tkey1].id)
                            data_count_z_list.append(commodity_data[tkey1].count)
                            data_price_z_list.append(commodity_data[tkey1].price)
                            data_sku_z_list.append(commodity_data[tkey2].id)
                            data_count_z_list.append(commodity_data[tkey2].count)
                            data_price_z_list.append(commodity_data[tkey2].price)
                        else:
                            data_sku_c_list.append(commodity_data[tkey1].id)
                            data_count_c_list.append(commodity_data[tkey1].count)
                            data_price_c_list.append(commodity_data[tkey1].price)
                            data_sku_c_list.append(commodity_data[tkey2].id)
                            data_count_c_list.append(commodity_data[tkey2].count)
                            data_price_c_list.append(commodity_data[tkey2].price)
                    if '米油杂粮' in it_di:
                        tkey1 = '金龙鱼雪域稻香贡 五常大米5kg地标产品0'
                        tkey2 = '千岁好有机山茶油礼盒500ml*2瓶双有机认证0'
                        tkey3 = '燕之坊 合家欢礼盒（9品12杂粮）3kg\n0'
                        commodity_data[tkey1].addcount(2)
                        commodity_data[tkey2].addcount(1)
                        commodity_data[tkey3].addcount(1)
                        data_sku_c_list.append(commodity_data[tkey1].id)
                        data_count_c_list.append(2)
                        data_price_c_list.append(commodity_data[tkey1].price)
                        data_sku_c_list.append(commodity_data[tkey2].id)
                        data_count_c_list.append(commodity_data[tkey2].count)
                        data_price_c_list.append(commodity_data[tkey2].price)
                        data_sku_c_list.append(commodity_data[tkey3].id)
                        data_count_c_list.append(commodity_data[tkey3].count)
                        data_price_c_list.append(commodity_data[tkey3].price)

        type = 0
        yanzhengall = 0
        for t_sku_list, t_count_list, t_price_list in zip([data_sku_z_list, data_sku_c_list], [data_count_z_list, data_count_c_list], [data_price_z_list, data_price_c_list]):
            if len(t_sku_list) == 0:
                type += 1
                continue
            data_price_all = 0
            for tprice, tcount in zip(t_price_list, t_count_list):
                data_price_all = data_price_all + tcount * tprice
            data_sku = '|'.join(t_sku_list)
            data_price = '|'.join(list(map(str, t_price_list)))
            data_count = '|'.join(list(map(str, t_count_list)))
            if '跳过' in data_saddr:
                data_sheng, data_shi, data_qv = '', '', ''
            else:
                try:
                    data_sheng, data_shi, data_qv = data_saddr.split('-')
                except:
                    # data_sheng, data_shi, data_qv = '其他地区', '其他地区', '其他地区'

                    tres = address_rec(data_saddr_all)
                    ts = list(tres.values())
                    tname = ts[6] if len(ts[6]) > 0 else ts[1][-3:0]
                    data_sheng, data_shi, data_qv = ts[5], ts[11], ts[10]
                    data_sname, data_sipone = tname, ts[3]


                    # writeData = [data_id, data_name, data_ipone, data_sku, data_price, data_count, data_price_all, data_price_all,
            #              data_sname, data_sipone, data_sheng, data_shi, data_qv, data_saddr_all]

            # writeData = [data_sku, data_price, data_count, data_price_all, data_price_all,
            #              data_sname, '', data_sipone, data_sheng, data_shi, data_qv, '', data_saddr_all, '', '', '']

            writeData = [data_sku, data_price, data_count, data_price_all,
                         data_sname, '', data_sipone, data_sheng, data_shi, data_qv, '', data_saddr_all, '', '', '']

            print(writeData)

            yanzhengall = yanzhengall + data_price_all
            if type:
                ws1.append(
                    writeData)
            else:
                ws2.append(
                    writeData)
            type += 1

        totalAcountprice[data_sname] = yanzhengall
        print('all price : {}'.format(yanzhengall))

    print(totalAcountprice)
    print('total acount all price : {}'.format(sum(totalAcountprice.values())))

    for i in commodity_data.keys():
        if commodity_data[i].total_count > 0:
            writeData = [commodity_data[i].id, commodity_data[i].name, commodity_data[i].total_count,
                         commodity_data[i].price, commodity_data[i].total_count*commodity_data[i].price*commodity_data[i].count,
                         commodity_data[i].geiyuangong, commodity_data[i].total_count*commodity_data[i].geiyuangong*commodity_data[i].count,
                         commodity_data[i].geishitang, commodity_data[i].total_count*commodity_data[i].geishitang*commodity_data[i].count]
            print(writeData)
            ws3.append(
                writeData)
    # 保存 Excel 文件
    wb.save(resultname + str(time.time()) + '.xlsx')


if __name__ == '__main__':
    info = input("请将所有文件放在同一目录，导出文件将在同一目录下生成。确认按回车")
    # i1_input = input("请输入商品文件名(不带文件名后缀 xxx.xlsx): ")
    # i2_input = input("请输入客户文件名(不带文件名后缀 xxx.xlsx): ")
    # o1_output = input("请输入导出文件名（不用后缀 xxx）: ")

    info = input("请将所有文件放在同一目录，导出文件将在同一目录下生成。确认按回车")
    i1_input = 'd2'
    i2_input = 's2'
    o1_output = 'r2'

    print("商品文件名: ", i1_input)
    print("客户文件名: ", i2_input)
    print("导出文件名: ", o1_output)
    root_path = './data/yyw/'
    foo(root_path + i1_input + '.xlsx', root_path + i2_input + '.xlsx', root_path + o1_output)